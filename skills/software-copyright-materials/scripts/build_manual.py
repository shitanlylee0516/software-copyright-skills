#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_manual.py

Generates an editable Word (.docx) software manual template for Chinese software copyright applications.
Parses Markdown, embeds real placeholder PNG slots for screenshots, applies Song/SimSun formatting,
sets headers/footers with PAGE/NUMPAGES, and marks unverified metadata as "待确认".
"""

import argparse
import io
import json
import os
import sys
from pathlib import Path

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    sys.stderr.write("Error: Pillow is required. Install via `pip install -r requirements-optional.txt`\n")
    sys.exit(1)

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    sys.stderr.write("Error: python-docx is required. Install via `pip install -r requirements-optional.txt`\n")
    sys.exit(1)

try:
    from markdown_it import MarkdownIt
except ImportError:
    sys.stderr.write("Error: markdown-it-py is required. Install via `pip install -r requirements-optional.txt`\n")
    sys.exit(1)


def create_placeholder_png(width_px: int, height_px: int, slot_label: str) -> io.BytesIO:
    """Creates a real in-memory grayscale PNG image for replacement in Word."""
    img = Image.new("L", (width_px, height_px), color=240)
    draw = ImageDraw.Draw(img)

    # Outer border
    draw.rectangle([(0, 0), (width_px - 1, height_px - 1)], outline=160, width=2)
    # Inner border
    draw.rectangle([(6, 6), (width_px - 7, height_px - 7)], outline=200, width=1)

    # Diagonal dashed/dotted lines
    step = 20
    for offset in range(-height_px, width_px, step):
        x1 = max(0, offset)
        y1 = max(0, -offset)
        x2 = min(width_px, offset + height_px)
        y2 = min(height_px, height_px - (offset + height_px - x2))
        draw.line([(x1, y1), (x2, y2)], fill=228, width=1)

    # Center card
    card_w = min(width_px - 40, 480)
    card_h = 130
    cx = width_px // 2
    cy = height_px // 2
    x0 = cx - card_w // 2
    y0 = cy - card_h // 2
    draw.rectangle([(x0, y0), (x0 + card_w, y0 + card_h)], fill=255, outline=140, width=2)

    font = ImageFont.load_default()

    line1 = f"[{slot_label}]"
    line2 = "Word: Right-click -> Change Picture"
    line3 = "(Replace with actual software UI screenshot)"

    def draw_centered_text(y_pos: int, text: str, fill: int = 40):
        bbox = draw.textbbox((0, 0), text, font=font)
        w = bbox[2] - bbox[0]
        draw.text((cx - w // 2, y_pos), text, font=font, fill=fill)

    draw_centered_text(y0 + 20, line1, fill=0)
    draw_centered_text(y0 + 55, line2, fill=60)
    draw_centered_text(y0 + 85, line3, fill=100)

    bio = io.BytesIO()
    img.save(bio, format="PNG")
    bio.seek(0)
    return bio


def add_field(p, field_name: str) -> None:
    """Adds a dynamic Word XML field (e.g., PAGE or NUMPAGES) to a paragraph."""
    run = p.add_run()
    r = run._r

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = f" {field_name} "

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def set_run_font(run, font_name: str = "SimSun", size_pt: float = 10.5, bold: bool = False, color_rgb: RGBColor = RGBColor(0, 0, 0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = color_rgb
    # Enable SimSun for East Asian characters
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)


def configure_document_styles(doc: Document) -> None:
    """Configures default styles (Normal, Title, Heading 1, 2, 3) to use SimSun/eastAsia, black, and title no border."""
    styles = doc.styles

    style_configs = {
        'Normal': (10.5, False),
        'Title': (20.0, True),
        'Heading 1': (15.0, True),
        'Heading 2': (14.0, True),
        'Heading 3': (12.0, True),
    }

    for style_name, (size_pt, is_bold) in style_configs.items():
        if style_name in styles:
            s = styles[style_name]
            s.font.name = "SimSun"
            s.font.size = Pt(size_pt)
            s.font.bold = is_bold
            s.font.color.rgb = RGBColor(0, 0, 0)

            # Ensure East Asian fonts and XML properties
            rPr = s._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.append(rFonts)
            rFonts.set(qn('w:ascii'), "SimSun")
            rFonts.set(qn('w:hAnsi'), "SimSun")
            rFonts.set(qn('w:eastAsia'), "SimSun")

            if style_name == 'Title':
                pPr = s._element.get_or_add_pPr()
                # Clear paragraph borders if any
                pBdr = pPr.find(qn('w:pBdr'))
                if pBdr is not None:
                    pPr.remove(pBdr)
                pBdr_none = parse_xml(
                    f'<w:pBdr {nsdecls("w")}>'
                    f'  <w:top w:val="none"/>'
                    f'  <w:left w:val="none"/>'
                    f'  <w:bottom w:val="none"/>'
                    f'  <w:right w:val="none"/>'
                    f'</w:pBdr>'
                )
                pPr.append(pBdr_none)


def style_table(table, col_widths=None) -> None:
    """Applies D9D9D9 borders, tblCellMar 100 twips, cantSplit, tblHeader, F5F5F5 header, and column widths."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr

    # Remove existing borders or cell margins if present
    for child in list(tblPr):
        if child.tag in (qn('w:tblBorders'), qn('w:tblCellMar')):
            tblPr.remove(child)

    # Set borders: all D9D9D9
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

    # Set cell margin (padding): 100 twips
    mar_xml = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'  <w:top w:w="100" w:type="dxa"/>'
        f'  <w:left w:w="100" w:type="dxa"/>'
        f'  <w:bottom w:w="100" w:type="dxa"/>'
        f'  <w:right w:w="100" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tblPr.append(mar_xml)

    # Apply column widths to grid if provided
    if col_widths:
        tblGrid = table._tbl.tblGrid
        if tblGrid is not None:
            # remove old tblGrid and rebuild
            table._tbl.remove(tblGrid)
        new_tblGrid = OxmlElement('w:tblGrid')
        for w in col_widths:
            dxa_val = int(w.twips)
            gridCol = OxmlElement('w:gridCol')
            gridCol.set(qn('w:w'), str(dxa_val))
            new_tblGrid.append(gridCol)
        # insert tblGrid after tblPr
        table._tbl.insert(table._tbl.index(tblPr) + 1, new_tblGrid)

    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>'))

        if col_widths:
            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(col_widths):
                    w = col_widths[c_idx]
                    dxa_val = int(w.twips)
                    cell.width = w
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcW = tcPr.find(qn('w:tcW'))
                    if tcW is None:
                        tcW = OxmlElement('w:tcW')
                        tcPr.append(tcW)
                    tcW.set(qn('w:w'), str(dxa_val))
                    tcW.set(qn('w:type'), 'dxa')


def render_inline_tokens(p, inline_tokens, default_size=10.5, base_bold=False):
    """Renders inline tokens (text, strong, code_inline, etc.) cleanly into a paragraph."""
    is_bold = base_bold
    for child in inline_tokens:
        if child.type == "text":
            run = p.add_run(child.content)
            set_run_font(run, "SimSun", default_size, bold=is_bold)
        elif child.type == "strong_open":
            is_bold = True
        elif child.type == "strong_close":
            is_bold = base_bold
        elif child.type == "code_inline":
            run = p.add_run(child.content)
            set_run_font(run, "Consolas", default_size - 0.5, bold=is_bold)
        elif child.type in ("softbreak", "hardbreak"):
            p.add_run("\n")


def parse_and_append_markdown(doc, md_text: str):
    md = MarkdownIt("commonmark", {"breaks": True}).enable("table")
    tokens = md.parse(md_text)

    i = 0
    num_tokens = len(tokens)
    list_stack = []

    while i < num_tokens:
        token = tokens[i]

        if token.type == "heading_open":
            level = int(token.tag[1])
            inline_tok = tokens[i + 1]

            # Use Word Heading styles: Heading 1, Heading 2, Heading 3
            style_name = f'Heading {level}' if level in (1, 2, 3) else 'Normal'
            p = doc.add_paragraph(style=style_name)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            font_size = 15 if level == 1 else (14 if level == 2 else 12)
            render_inline_tokens(p, inline_tok.children or [], default_size=font_size, base_bold=True)
            i += 3  # heading_open, inline, heading_close
            continue

        elif token.type == "paragraph_open":
            inline_tok = tokens[i + 1]
            if list_stack:
                cur = list_stack[-1]
                nesting_level = len(list_stack)
                if cur["type"] == "ordered":
                    p = doc.add_paragraph(style='Normal')
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.25
                    base_indent = Inches(0.25 * nesting_level)
                    p.paragraph_format.left_indent = base_indent
                    p.paragraph_format.first_line_indent = Inches(-0.25)
                    if cur["pending_prefix"]:
                        prefix = f"{cur['current_num']}. "
                        cur["pending_prefix"] = False
                        run = p.add_run(prefix)
                        set_run_font(run, "SimSun", 10.5, bold=False)
                else:
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    if nesting_level > 1:
                        p.paragraph_format.left_indent = Inches(0.25 * (nesting_level - 1))
            else:
                p = doc.add_paragraph(style='Normal')
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.25

            render_inline_tokens(p, inline_tok.children or [], default_size=10.5, base_bold=False)
            i += 3  # paragraph_open, inline, paragraph_close
            continue

        elif token.type == "bullet_list_open":
            list_stack.append({"type": "bullet"})
            i += 1
            continue

        elif token.type == "bullet_list_close":
            if list_stack:
                list_stack.pop()
            i += 1
            continue

        elif token.type == "ordered_list_open":
            start_val = 1
            if hasattr(token, 'attrs') and token.attrs:
                attrs_dict = dict(token.attrs) if isinstance(token.attrs, (list, tuple)) else token.attrs
                if "start" in attrs_dict:
                    try:
                        start_val = int(attrs_dict["start"])
                    except (ValueError, TypeError):
                        start_val = 1
            list_stack.append({
                "type": "ordered",
                "next_num": start_val,
                "current_num": start_val,
                "pending_prefix": False,
            })
            i += 1
            continue

        elif token.type == "ordered_list_close":
            if list_stack:
                list_stack.pop()
            i += 1
            continue

        elif token.type == "list_item_open":
            if list_stack and list_stack[-1]["type"] == "ordered":
                cur = list_stack[-1]
                cur["current_num"] = cur["next_num"]
                cur["next_num"] += 1
                cur["pending_prefix"] = True
            i += 1
            continue

        elif token.type == "list_item_close":
            i += 1
            continue

        elif token.type == "table_open":
            # Collect table structure
            table_tokens = []
            j = i
            while j < num_tokens and tokens[j].type != "table_close":
                table_tokens.append(tokens[j])
                j += 1
            table_tokens.append(tokens[j])  # include table_close
            i = j + 1

            # Extract rows with parsed inline tokens
            rows_cells = []
            current_row = []
            for t in table_tokens:
                if t.type == "tr_open":
                    current_row = []
                elif t.type == "tr_close":
                    rows_cells.append(current_row)
                elif t.type == "inline":
                    current_row.append(t.children or [])

            if rows_cells:
                num_cols = len(rows_cells[0])
                table = doc.add_table(rows=len(rows_cells), cols=num_cols)

                # Assign default proportional column widths fitting A4 text width (16.6 cm)
                if num_cols == 2:
                    col_widths = [Cm(4.2), Cm(12.4)]
                elif num_cols == 4:
                    col_widths = [Cm(1.3), Cm(4.0), Cm(3.0), Cm(8.3)]
                else:
                    per_col = 16.6 / num_cols
                    col_widths = [Cm(per_col)] * num_cols

                style_table(table, col_widths=col_widths)

                for r_idx, row in enumerate(rows_cells):
                    for c_idx, cell_inlines in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                        render_inline_tokens(p, cell_inlines, default_size=10, base_bold=(r_idx == 0))
            continue

        elif token.type in ("fence", "code_block"):
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(token.content)
            set_run_font(run, "Consolas", 9.5)
            i += 1
            continue

        else:
            i += 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Build editable DOCX software manual template.")
    parser.add_argument("--input", required=True, help="Path to input Markdown file.")
    parser.add_argument("--metadata", required=True, help="Path to metadata JSON file.")
    parser.add_argument("--screenshots", required=True, help="Path to screenshots JSON file.")
    parser.add_argument("--output", required=True, help="Path to output .docx file.")
    args = parser.parse_args()

    input_path = Path(args.input)
    meta_path = Path(args.metadata)
    shots_path = Path(args.screenshots)
    out_path = Path(args.output)

    # 1. Output file must not already exist
    if out_path.exists():
        sys.stderr.write(f"Error: Output file already exists: {out_path.name}\n")
        sys.exit(1)

    # 2. Check input existence and read
    if not input_path.is_file():
        sys.stderr.write(f"Error: Markdown input file not found: {input_path.name}\n")
        sys.exit(1)
    if not meta_path.is_file():
        sys.stderr.write(f"Error: Metadata JSON file not found: {meta_path.name}\n")
        sys.exit(1)
    if not shots_path.is_file():
        sys.stderr.write(f"Error: Screenshots JSON file not found: {shots_path.name}\n")
        sys.exit(1)

    # Read strictly without mutating input files, with UTF-8 BOM support
    raw_md_bytes = input_path.read_bytes()
    md_text = raw_md_bytes.decode("utf-8-sig")

    try:
        raw_meta = meta_path.read_bytes().decode("utf-8-sig")
        meta_data = json.loads(raw_meta)
    except Exception as e:
        sys.stderr.write(f"Error reading metadata JSON: {e}\n")
        sys.exit(1)

    if not isinstance(meta_data, dict):
        sys.stderr.write("Error: Metadata must be a JSON object.\n")
        sys.exit(1)

    try:
        raw_shots = shots_path.read_bytes().decode("utf-8-sig")
        screenshots_data = json.loads(raw_shots)
    except Exception as e:
        sys.stderr.write(f"Error reading screenshots JSON: {e}\n")
        sys.exit(1)

    # Screenshots must be a list
    if not isinstance(screenshots_data, list):
        sys.stderr.write("Error: Screenshots configuration must be a JSON list.\n")
        sys.exit(1)

    # Validate each screenshot item strictly: id, caption strings, aspect portrait|landscape
    for idx, shot in enumerate(screenshots_data, start=1):
        if not isinstance(shot, dict):
            sys.stderr.write(f"Error: Screenshot entry {idx} must be a JSON object.\n")
            sys.exit(1)
        if "id" not in shot or not isinstance(shot["id"], str) or not shot["id"].strip():
            sys.stderr.write(f"Error: Screenshot entry {idx} must have non-empty string 'id'.\n")
            sys.exit(1)
        if "caption" not in shot or not isinstance(shot["caption"], str) or not shot["caption"].strip():
            sys.stderr.write(f"Error: Screenshot entry {idx} must have non-empty string 'caption'.\n")
            sys.exit(1)
        if "aspect" not in shot or not isinstance(shot["aspect"], str):
            sys.stderr.write(f"Error: Screenshot entry {idx} missing 'aspect'. Must be 'portrait' or 'landscape'.\n")
            sys.exit(1)
        aspect_val = shot["aspect"].strip().lower()
        if aspect_val not in ("portrait", "landscape"):
            sys.stderr.write(f"Error: Screenshot entry {idx} has invalid aspect '{shot['aspect']}'. Must be 'portrait' or 'landscape'.\n")
            sys.exit(1)

    # Extract metadata with "待确认" fallback
    def get_meta(key: str, default: str = "待确认") -> str:
        val = meta_data.get(key)
        if val is None or str(val).strip() == "":
            return default
        return str(val).strip()

    software_name = get_meta("name", "待确认软件名称")
    software_version = get_meta("version", "待确认版本")
    owner_name = get_meta("owner", "待确认著作权人")
    completion_date = get_meta("completion_date", "待确认开发完成日期")
    publication_date = get_meta("publication_date", "待确认首次发表日期")

    # Initialize Word Document
    doc = Document()
    configure_document_styles(doc)

    # Configure Section: A4, 2.2 cm margins (width=21.0, margins 2.2+2.2 leaves 16.6 cm text area)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # Header: real name / version, black 9pt
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header_p.add_run(f"{software_name} {software_version} 用户操作手册")
    set_run_font(hrun, "SimSun", 9.0, color_rgb=RGBColor(0, 0, 0))

    # Footer: PAGE / NUMPAGES, gray 9pt
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun1 = footer_p.add_run("第 ")
    set_run_font(frun1, "SimSun", 9.0, color_rgb=RGBColor(128, 128, 128))
    add_field(footer_p, "PAGE")
    frun2 = footer_p.add_run(" 页 / 共 ")
    set_run_font(frun2, "SimSun", 9.0, color_rgb=RGBColor(128, 128, 128))
    add_field(footer_p, "NUMPAGES")
    frun3 = footer_p.add_run(" 页")
    set_run_font(frun3, "SimSun", 9.0, color_rgb=RGBColor(128, 128, 128))

    # Title & Metadata block
    title_p = doc.add_paragraph(style='Title')
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(12)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title_p.add_run(f"{software_name} {software_version}\n用户操作手册（待事实核验草稿）")
    set_run_font(trun, "SimSun", 20.0, bold=True, color_rgb=RGBColor(0, 0, 0))

    # Metadata summary table (4.2cm / 12.4cm)
    meta_table = doc.add_table(rows=5, cols=2)
    style_table(meta_table, col_widths=[Cm(4.2), Cm(12.4)])
    meta_fields = [
        ("软件名称与版本", f"{software_name} {software_version}"),
        ("软件著作权人", owner_name),
        ("开发完成日期", completion_date),
        ("首次发表日期", publication_date),
        ("文档审查状态", "草稿：待补充真实运行截图并核验所有主体与日期事实"),
    ]
    for r_idx, (k, v) in enumerate(meta_fields):
        c0 = meta_table.cell(r_idx, 0)
        c1 = meta_table.cell(r_idx, 1)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r1 = p1.add_run(v)
        set_run_font(r0, "SimSun", 10.0, bold=True)
        set_run_font(r1, "SimSun", 10.0, bold=False)

    doc.add_paragraph(style='Normal')  # Spacing

    # Parse and render Markdown body
    parse_and_append_markdown(doc, md_text)

    # Append Screenshots section at the end
    if screenshots_data:
        doc.add_page_break()
        sec_p = doc.add_paragraph(style='Heading 1')
        sec_p.paragraph_format.space_before = Pt(12)
        sec_p.paragraph_format.space_after = Pt(8)
        srun = sec_p.add_run("软件界面截图附录（占位槽位与替换说明）")
        set_run_font(srun, "SimSun", 15.0, bold=True)

        guide_p = doc.add_paragraph(style='Normal')
        guide_p.paragraph_format.space_before = Pt(2)
        guide_p.paragraph_format.space_after = Pt(8)
        grun = guide_p.add_run(
            "说明：以下各槽位包含真实嵌入的占位图片。请在 Word 中右键单击占位图片，选择【更改图片】（Change Picture）"
            "直接替换为实际运行界面全屏或主窗口截图。请勿直接删除再插入，以保留预设版心与版面布局。"
        )
        set_run_font(grun, "SimSun", 10.0, color_rgb=RGBColor(80, 80, 80))

        # Screenshot summary table (4 columns: 1.3cm / 4.0cm / 3.0cm / 8.3cm = 16.6cm)
        shots_table = doc.add_table(rows=len(screenshots_data) + 1, cols=4)
        style_table(shots_table, col_widths=[Cm(1.3), Cm(4.0), Cm(3.0), Cm(8.3)])

        headers = ["序号", "图号标识", "版式方向", "说明 / 替换提示"]
        for c_idx, h in enumerate(headers):
            cell = shots_table.cell(0, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h)
            set_run_font(run, "SimSun", 10.0, bold=True)

        for s_idx, shot in enumerate(screenshots_data, start=1):
            s_id = shot["id"].strip()
            s_cap = shot["caption"].strip()
            s_aspect = shot["aspect"].strip().lower()
            aspect_label = "竖屏 (9:16)" if s_aspect == "portrait" else "横屏 (16:9)"

            row_cells = [str(s_idx), s_id, aspect_label, s_cap]
            for c_idx, val in enumerate(row_cells):
                cell = shots_table.cell(s_idx, c_idx)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                set_run_font(run, "SimSun", 9.5, bold=False)

        for idx, shot in enumerate(screenshots_data, start=1):
            slot_id = shot["id"].strip()
            caption = shot["caption"].strip()
            aspect = shot["aspect"].strip().lower()

            doc.add_page_break()

            # Instruction paragraph outside image
            inst_p = doc.add_paragraph(style='Normal')
            inst_p.paragraph_format.space_before = Pt(6)
            inst_p.paragraph_format.space_after = Pt(4)
            inst_p.paragraph_format.keep_with_next = True
            irun = inst_p.add_run(f"【槽位 {idx} 说明】右键图片 -> 更改图片 -> 替换为「{caption}」的实际运行截图")
            set_run_font(irun, "SimSun", 9.5, color_rgb=RGBColor(100, 100, 100))

            # Image paragraph
            img_p = doc.add_paragraph(style='Normal')
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_before = Pt(4)
            img_p.paragraph_format.space_after = Pt(4)
            img_p.paragraph_format.keep_with_next = True

            if aspect == "portrait":
                # 9:16 aspect ratio: 7.4 cm x 13.1556 cm
                w_cm = Cm(7.4)
                h_cm = Cm(13.1556)
                px_w, px_h = 450, 800
            else:
                # 16:9 aspect ratio: 15 cm x 8.4375 cm
                w_cm = Cm(15.0)
                h_cm = Cm(8.4375)
                px_w, px_h = 800, 450

            # ASCII-only label for PIL image
            slot_label = f"SCREENSHOT {idx}"
            png_bio = create_placeholder_png(px_w, px_h, slot_label)
            shape = img_p.add_run().add_picture(png_bio, width=w_cm, height=h_cm)

            # Set shape._inline.docPr descr to id+caption+replacement guidance
            descr_text = f"{slot_id} {caption} - 请在 Word 中右键单击占位图片选择「更改图片」替换为实际运行截图"
            shape._inline.docPr.set('descr', descr_text)

            # Caption paragraph
            cap_p = doc.add_paragraph(style='Normal')
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(4)
            cap_p.paragraph_format.space_after = Pt(12)
            crun = cap_p.add_run(f"{slot_id} {caption}")
            set_run_font(crun, "SimSun", 10.0, bold=True)

    # Save output document cleanly
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


if __name__ == "__main__":
    main()
